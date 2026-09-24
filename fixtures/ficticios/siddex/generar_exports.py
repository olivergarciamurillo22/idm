"""Genera los exports ficticios de Siddex en fixtures/ficticios/siddex/ (articulos, proveedores, fabricantes, pedidos, stock)
y un planning ficticio en fixtures/ficticios/planning_ficticio.xlsx con la misma estructura que el real (hoja, filas 5+, B..E).
Ejecutar: PYTHONPATH=src .venv/bin/python fixtures/ficticios/siddex/generar_exports.py"""

from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent


def hoja(ruta: Path, cabeceras: list[str], filas: list[list]) -> None:
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.append(["Export ficticio de Siddex"])
    ws.append([])
    ws.append(cabeceras)
    for fila in filas:
        ws.append(fila)
    libro.save(ruta)


def main() -> None:
    hoja(
        AQUI / "articulos.xlsx",
        ["Código", "Descripción", "Unidad", "Proveedor", "Precio", "Dto", "Múltiplo", "Tipo"],
        [
            ["T30.000.100", "TORNILLO M10x30 DIN933", "UD", 9003, 0.08, 0, 100, "Comercial"],
            ["R40.000.001", "RUEDA 400x8 LISA", "UD", 9001, 38.00, 0, 1, "Comercial"],
            ["H50.000.020", "BOMBA HIDRAULICA 20L", "UD", 9002, 210.00, 0, 1, "Comercial"],
            ["I37.000.012", "MANGUERA HIDRAULICA 1/2", "M", 9001, 6.40, 0, 50, "Comercial"],
            ["I37.000.099", "RACOR 1/2 MACHO", "UD", 9001, 1.10, 0, 10, "Comercial"],
            ["I34.000.842", "TERMINAL FASTON 9,4", "UD", 9003, 0.30, 45, 100, "Comercial"],
            ["I34.000.900", "RELE 24V 40A", "UD", 9003, 3.10, 10, 1, "Comercial"],
            ["134.000.842", "TERMINAL FASTON H. 9,4 AMAR", "UD", 9001, 0.30, 45, 100, "Comercial"],
            ["I33.000.786", "RODAMIENTO 6205 2RS", "UD", 9001, 9.14, 45, 1, "Comercial"],
            ["P20.000.010", "CORTE TUBO 80x80x3 L=1200", "UD", None, 12.50, 0, 1, "MP"],
            ["P20.000.011", "CHAPA KG 3MM", "KG", None, 1.20, 0, 1, "MP"],
            ["MO01001", "SERVICIO TURISMO (CONTRATOS)", "UD", 9002, None, 0, 1, "Comercial"],
        ],
    )
    hoja(
        AQUI / "proveedores.xlsx",
        ["Código", "Nombre", "CIF", "Email"],
        [
            [9001, "SUMINISTROS FICTICIOS LA VEGA S.L.", "B00000001", "pedidos@vega.ficticio"],
            [9002, "RECAMBIOS EJEMPLO S.A.", "A00000002", "pedidos@recambios.ficticio"],
            [9003, "ELECTRO FICTICIO S.L.", "B00000003", "ventas@electro.ficticio"],
        ],
    )
    hoja(
        AQUI / "fabricantes.xlsx",
        ["Código", "Proveedor", "Código alternativo", "Descripción proveedor"],
        [
            ["134.000.842", 9001, "PVCC550942", "TERMINAL FASTON H 9,4 AMARILLO"],
            ["I33.000.786", 9001, "ABR-6205", "RODAMIENTO 6205 2RS"],
            ["I37.000.012", 9001, "MNG-12", "MANGUERA HIDRAULICA 1/2"],
            ["MO01001", 9002, "MO01001", "SERVICIO TURISMO (CONTRATOS)"],
            ["I34.000.842", 9003, "EL-FAS-94", "TERMINAL FASTON 9,4"],
            ["I34.000.900", 9003, "EL-REL-24", "RELE 24V 40A"],
            ["T30.000.100", 9003, "EL-TOR-1030", "TORNILLO M10x30"],
        ],
    )
    hoja(
        AQUI / "pedidos.xlsx",
        [
            "Pedido",
            "Proveedor",
            "Fecha",
            "Artículo",
            "Descripción",
            "Cantidad",
            "Unidad",
            "Precio",
            "Dto",
            "Recibida",
            "Referencia proveedor",
        ],
        [
            [
                20261060,
                9001,
                "2026-09-10",
                "134.000.842",
                "TERMINAL FASTON H. 9,4 AMAR",
                100,
                "UD",
                0.30,
                45,
                0,
                "PVCC550942",
            ],
            [20261060, 9001, "2026-09-10", "I33.000.786", "RODAMIENTO 6205 2RS", 5, "UD", 9.14, 45, 0, "ABR-6205"],
            [20261060, 9001, "2026-09-10", "I37.000.012", "MANG. HIDR. 1/2", 5, "M", 6.40, 0, 0, "MNG-12"],
            [20261061, 9003, "2026-09-11", "I34.000.900", "RELE 24V 40A", 12, "UD", 3.10, 10, 0, "EL-REL-24"],
            [20261061, 9003, "2026-09-11", "T30.000.100", "TORNILLO M10x30", 100, "UD", 0.08, 0, 0, "EL-TOR-1030"],
            [20261055, 9002, "2026-09-01", "H50.000.020", "BOMBA HIDRAULICA 20L", 2, "UD", 210.00, 0, 2, None],
        ],
    )
    hoja(
        AQUI / "stock.xlsx",
        ["Código", "Stock"],
        [["R40.000.001", 2], ["H50.000.020", 1], ["T30.000.100", 500], ["I37.000.012", 0]],
    )

    planning = openpyxl.Workbook()
    ws = planning.active
    ws.title = "SEPT_26"
    ws["B2"] = "PLANNING FICTICIO"
    ws.append([])
    ws.append([None, "PRODUCTO", "CANTIDAD", "CLIENTE", "VENDEDOR"])
    for fila in [
        [None, "MAQUINA FICTICIA UNO", 2, "CLIENTE A", "V1"],
        [None, "MAQUINA FICTICIA UNO L", 1, "CLIENTE B", "V1"],
        [None, "MAQUINA FICTICIA UNO 2500-900 ED", 1, "CLIENTE C", "V2"],
        [None, "MAQUINA DESCONOCIDA", 1, "CLIENTE D", "V2"],
        [None, None, None, None, None],
    ]:
        ws.append(fila)
    planning.save(AQUI.parent / "planning_ficticio.xlsx")
    print("exports ficticios escritos en", AQUI)


if __name__ == "__main__":
    main()
