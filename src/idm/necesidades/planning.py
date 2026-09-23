"""Lee el Excel del planning mensual: una hoja por mes (ENERO_26…), cabecera PRODUCTO/CANTIDAD/CLIENTE/VENDEDOR.
Devuelve líneas con el nombre del producto en texto libre tal cual está, sin interpretar variantes.
No conoce Siddex ni los códigos: eso lo hace mapa.py."""

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import openpyxl

from idm.dominio.dinero import a_decimal


@dataclass
class LineaPlanning:
    producto: str
    cantidad: Decimal
    cliente: str
    vendedor: str
    fila: int


def hojas(ruta: Path) -> list[str]:
    libro = openpyxl.load_workbook(ruta, read_only=True)
    nombres = libro.sheetnames
    libro.close()
    return nombres


def leer_planning(ruta: Path, hoja: str) -> list[LineaPlanning]:
    """Busca la fila de cabecera (la que contiene PRODUCTO) y lee hasta la primera fila sin producto."""
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    if hoja not in libro.sheetnames:
        libro.close()
        raise ValueError(f"La hoja '{hoja}' no existe en {ruta.name}. Hojas: {', '.join(libro.sheetnames)}")
    ws = libro[hoja]
    lineas: list[LineaPlanning] = []
    col_producto = None
    for numero, fila in enumerate(ws.iter_rows(values_only=True), start=1):
        celdas = [str(c).strip() if c is not None else "" for c in fila]
        if col_producto is None:
            if any(c.upper() == "PRODUCTO" for c in celdas):
                col_producto = next(i for i, c in enumerate(celdas) if c.upper() == "PRODUCTO")
            continue
        producto = celdas[col_producto] if col_producto < len(celdas) else ""
        if not producto:
            if lineas:
                break
            continue
        cantidad = a_decimal(fila[col_producto + 1]) if col_producto + 1 < len(fila) else None
        lineas.append(
            LineaPlanning(
                producto=producto,
                cantidad=cantidad or Decimal("0"),
                cliente=celdas[col_producto + 2] if col_producto + 2 < len(celdas) else "",
                vendedor=celdas[col_producto + 3] if col_producto + 3 < len(celdas) else "",
                fila=numero,
            )
        )
    libro.close()
    if col_producto is None:
        raise ValueError(f"No se encontró la cabecera PRODUCTO en la hoja {hoja}")
    return lineas
