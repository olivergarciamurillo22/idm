"""Escribe el resultado del cálculo de necesidades en un Excel para que Fernandillo lo revise.
Hojas: Necesidades (una fila por artículo con máquinas en las que sale), Fuera circuito (hierro), Incidencias.
No decide nada: es solo la salida legible del cálculo."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from idm.necesidades.calculo import Necesidad, ResultadoNecesidades

CABECERAS = [
    "Código",
    "Descripción",
    "Unidad",
    "Tipo",
    "Proveedor",
    "Cantidad bruta",
    "Stock",
    "Pendiente recibir",
    "Neta",
    "Múltiplo",
    "A pedir",
    "Precio",
    "Importe",
    "Máquinas",
]


def _fila(n: Necesidad) -> list:
    maquinas = ", ".join(f"{m}×{c}" for m, c in n.maquinas.items())
    return [
        n.codigo,
        n.descripcion,
        n.unidad,
        n.tipo,
        n.proveedor or "",
        float(n.cantidad_bruta),
        float(n.stock),
        float(n.pendiente_recibir),
        float(n.cantidad_neta),
        float(n.multiplo),
        float(n.cantidad_pedir),
        float(n.precio) if n.precio is not None else None,
        float(n.importe),
        maquinas,
    ]


def escribir_excel(resultado: ResultadoNecesidades, ruta: Path) -> Path:
    libro = Workbook()
    ws = libro.active
    ws.title = "Necesidades"
    ws.append(
        [
            f"Necesidades {resultado.hoja}",
            "",
            "",
            "Máquinas:",
            ", ".join(f"{m}×{c}" for m, c in resultado.maquinas.items()),
        ]
    )
    ws.append(CABECERAS)
    for celda in ws[2]:
        celda.font = Font(bold=True)
    for n in resultado.necesidades:
        ws.append(_fila(n))
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", "", "", "", "", "", float(resultado.total)])

    ws2 = libro.create_sheet("Fuera circuito")
    ws2.append(["Hierro y otros fuera del circuito (se piden por WhatsApp, sin pedido automático)"])
    ws2.append(CABECERAS)
    for n in resultado.fuera_circuito:
        ws2.append(_fila(n))

    ws3 = libro.create_sheet("Incidencias")
    ws3.append(["Incidencia"])
    for texto in resultado.incidencias:
        ws3.append([texto])

    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta
