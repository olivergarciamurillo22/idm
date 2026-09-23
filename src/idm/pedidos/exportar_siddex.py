"""Hoja Excel con las líneas de cada pedido en el orden de la pantalla de Siddex, para importar o teclear en un minuto.
Una hoja por pedido: proveedor (código Siddex), nuestro pedido, artículo, descripción, cantidad, ud., precio, dto.
Si Siddex admite importar desde fichero (pregunta pendiente), este fichero será el que se importe."""

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from idm.dominio.modelos import Pedido, Proveedor

CABECERAS = [
    "Proveedor (código Siddex)",
    "Nuestro pedido",
    "Fecha",
    "Artículo",
    "Descripción",
    "Cantidad",
    "Unidad",
    "Precio",
    "Dto %",
]


def titulo_hoja(numero: str) -> str:
    """Excel no admite / \\ ? * [ ] : en el nombre de hoja ni más de 31 caracteres."""
    return re.sub(r"[/\\?*\[\]:]", "-", numero)[:31]


def exportar(pedidos: list[Pedido], proveedores: dict[str, Proveedor], ruta: Path) -> Path:
    libro = Workbook()
    libro.remove(libro.active)
    for pedido in pedidos:
        ws = libro.create_sheet(titulo_hoja(pedido.numero))
        ws.append(CABECERAS)
        for c in ws[1]:
            c.font = Font(bold=True)
        prov = proveedores.get(pedido.proveedor)
        for li in pedido.lineas:
            ws.append(
                [
                    prov.codigo_siddex if prov else pedido.proveedor,
                    pedido.numero,
                    pedido.fecha.isoformat() if pedido.fecha else "",
                    li.codigo_idm,
                    li.descripcion,
                    float(li.cantidad),
                    li.unidad,
                    float(li.precio_bruto),
                    float(li.descuento_pct),
                ]
            )
    if not libro.sheetnames:
        libro.create_sheet("Sin pedidos").append(["No hay pedidos que exportar"])
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta
