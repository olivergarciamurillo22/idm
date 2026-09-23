"""Tests de necesidades/: planning, mapa con reglas L/E, cálculo agregado, múltiplos, stock, Excel de salida.
Usa el planning y los exports ficticios de fixtures/.
No toca datos/."""

from decimal import Decimal

import openpyxl

from idm.necesidades import mapa as mapa_mod
from idm.necesidades.calculo import calcular, redondear_a_multiplo
from idm.necesidades.planning import hojas, leer_planning
from idm.necesidades.salida import escribir_excel
from idm.necesidades.stock import descontar
from idm.siddex.desde_excel import SiddexDesdeExcel


def test_leer_planning(fixtures):
    ruta = fixtures / "planning_ficticio.xlsx"
    assert hojas(ruta) == ["SEPT_26"]
    lineas = leer_planning(ruta, "SEPT_26")
    assert [(li.producto, li.cantidad) for li in lineas] == [
        ("MAQUINA FICTICIA UNO", Decimal("2")),
        ("MAQUINA FICTICIA UNO L", Decimal("1")),
        ("MAQUINA FICTICIA UNO 2500-900 ED", Decimal("1")),
        ("MAQUINA DESCONOCIDA", Decimal("1")),
    ]
    assert lineas[0].cliente == "CLIENTE A"


def test_variante_desde_tokens():
    assert mapa_mod.variante_desde_tokens(["2250"]) == "base"
    assert mapa_mod.variante_desde_tokens(["L"]) == "L"
    assert mapa_mod.variante_desde_tokens(["ED"]) == "E"
    assert mapa_mod.variante_desde_tokens(["2500", "900", "E", "L"]) == "LE"
    assert mapa_mod.variante_desde_tokens(["CON", "ESCALON"]) == "base"


def test_resolver_mapa_real_a4r():
    mapa = mapa_mod.cargar_mapa()
    assert mapa_mod.resolver("A4R-2000-820", mapa).codigo == "I91.115.000"
    assert mapa_mod.resolver("A4R 2250 CON ESCALÓN", mapa).codigo == "I91.115.000"
    r = mapa_mod.resolver("A4R L 2250", mapa)
    assert r.variante == "L" and r.codigo is None and r.metodo == "sin_codigo"
    assert mapa_mod.resolver("A4R ED", mapa).variante == "E"
    assert mapa_mod.resolver("FUMIMATIC RUEDAS AGRICOLAS", mapa).codigo == "I91.640.000"
    assert mapa_mod.resolver("COSA RARA", mapa).metodo == "no_mapeado"


def test_redondear_a_multiplo():
    assert redondear_a_multiplo(Decimal("13"), Decimal("10")) == Decimal("20")
    assert redondear_a_multiplo(Decimal("20"), Decimal("10")) == Decimal("20")
    assert redondear_a_multiplo(Decimal("0"), Decimal("10")) == Decimal("0")
    assert redondear_a_multiplo(Decimal("2.5"), Decimal("1")) == Decimal("3")


def test_calcular_agrega_por_articulo_y_maquina(fixtures):
    gateway = SiddexDesdeExcel(fixtures / "siddex")
    mapa = mapa_mod.cargar_mapa(fixtures / "siddex" / "mapa_productos_ficticio.csv")
    planning = leer_planning(fixtures / "planning_ficticio.xlsx", "SEPT_26")
    r = calcular(planning, gateway, mapa, hoja="SEPT_26")
    # La variante E ("... 2500-900 ED") no tiene código todavía: queda como incidencia, no se calcula.
    assert r.maquinas == {"M91.000.001": Decimal("2"), "M91.000.002": Decimal("1")}
    por = {n.codigo: n for n in r.necesidades}
    # Tornillo: 48/máquina × 2 base + 40 × 1 L = 136, múltiplo 100 → 200
    assert por["T30.000.100"].cantidad_bruta == Decimal("136")
    assert por["T30.000.100"].cantidad_pedir == Decimal("200")
    assert por["T30.000.100"].maquinas == {"M91.000.001": Decimal("2"), "M91.000.002": Decimal("1")}
    assert por["T30.000.100"].proveedor == "FICT_ELECTRO"
    # Mangueras por rollos de 50: 4 × 2 = 8 → 50
    assert por["I37.000.012"].cantidad_pedir == Decimal("50")
    assert por["I37.000.012"].importe == Decimal("320.00")
    # Rueda lisa solo en la L
    assert por["R40.000.002"].maquinas == {"M91.000.002": Decimal("1")}
    # Hierro aparte
    assert {n.codigo for n in r.fuera_circuito} == {"P20.000.010", "P20.000.011"}
    assert "P20.000.010" not in por
    # Relé sin precio en el escandallo: lo toma del maestro
    assert por["I34.000.900"].precio == Decimal("3.1")
    # Incidencias: máquina desconocida y variante E sin código
    assert any("MAQUINA DESCONOCIDA" in t for t in r.incidencias)
    assert any("ED" in t and "variante E" in t for t in r.incidencias)


def test_descontar_stock_solo_en_codigos_fiables(fixtures):
    gateway = SiddexDesdeExcel(fixtures / "siddex")
    mapa = mapa_mod.cargar_mapa(fixtures / "siddex" / "mapa_productos_ficticio.csv")
    planning = leer_planning(fixtures / "planning_ficticio.xlsx", "SEPT_26")
    r = calcular(planning, gateway, mapa)
    stock, pendiente = gateway.stock(), gateway.pendiente_recibir()
    descontar(r.necesidades, stock, pendiente, codigos_fiables=frozenset({"H50.000.020"}))
    por = {n.codigo: n for n in r.necesidades}
    assert por["H50.000.020"].cantidad_bruta == Decimal("3")
    assert por["H50.000.020"].stock == Decimal("1")
    assert por["H50.000.020"].cantidad_pedir == Decimal("2")
    assert por["T30.000.100"].stock == Decimal("0")  # no fiable: no se descuenta
    descontar(r.necesidades, stock, pendiente, descontar_todo=True)
    assert por["T30.000.100"].cantidad_pedir == Decimal("0")  # 136 - 500 stock


def test_escribir_excel(fixtures, tmp_path):
    gateway = SiddexDesdeExcel(fixtures / "siddex")
    mapa = mapa_mod.cargar_mapa(fixtures / "siddex" / "mapa_productos_ficticio.csv")
    r = calcular(leer_planning(fixtures / "planning_ficticio.xlsx", "SEPT_26"), gateway, mapa, hoja="SEPT_26")
    ruta = escribir_excel(r, tmp_path / "necesidades.xlsx")
    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames == ["Necesidades", "Fuera circuito", "Incidencias"]
    filas = list(libro["Necesidades"].iter_rows(values_only=True))
    assert filas[1][0] == "Código"
    assert any(f[0] == "T30.000.100" and f[10] == 200 for f in filas)
